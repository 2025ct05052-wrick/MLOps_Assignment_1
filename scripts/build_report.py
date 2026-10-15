"""Build the ten-page MLOps assignment report from verified project outputs."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "MLOps_Assignment_Report.docx"
METRICS_PATH = ROOT / "artifacts" / "metrics.json"
EDA_PATH = ROOT / "artifacts" / "eda" / "eda_summary.json"
VERIFICATION_PATH = ROOT / "reports" / "verification_status.json"
NAVY = "17324D"
BLUE = "2F6690"
PALE = "EAF1F7"
GRAY = "53616F"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def style_run(run, size=10.5, color="263238", bold=False, name="Aptos") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_paragraph(doc, text="", size=10.5, color="263238", bold=False, before=0, after=6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    style_run(run, size=size, color=color, bold=bold)
    return paragraph


def add_bullet(doc, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.06
    style_run(paragraph.add_run(text), size=10.2)


def add_heading(doc, title: str, level=1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Heading 1" if level == 1 else "Heading 2"]
    paragraph.paragraph_format.space_before = Pt(5 if level == 1 else 2)
    paragraph.paragraph_format.space_after = Pt(7 if level == 1 else 4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title)
    style_run(run, size=17 if level == 1 else 12, color="000000", bold=True)


def add_kicker(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text.upper())
    style_run(run, size=8.5, color=BLUE, bold=True)


def add_table(doc, headers, rows, widths=None, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    table_properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table_properties.append(borders)
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(heading)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            style_run(run, size=font_size, color="FFFFFF", bold=True)
        if widths:
            cell.width = Inches(widths[index])
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            if row_number % 2 == 1:
                set_cell_shading(cells[index], "F3F6F9")
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    style_run(run, size=font_size, color="263238")
            if widths:
                cells[index].width = Inches(widths[index])
    return table


def add_page_break(doc):
    doc.add_page_break()


def add_page_title(doc, kicker: str, title: str):
    add_kicker(doc, kicker)
    add_heading(doc, title)


def make_architecture(path: Path) -> None:
    width, height = 1700, 410
    image = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 29)
        small = ImageFont.truetype("arial.ttf", 22)
    except OSError:
        font = ImageFont.load_default()
        small = ImageFont.load_default()
    cards = [
        ("UCI data", "303 rows\n13 predictors"),
        ("Data + EDA", "validation\nplots + summary"),
        ("CV training", "impute + encode\nLR + RF search"),
        ("MLflow", "params + metrics\nmodel + ROC plots"),
        ("FastAPI", "/predict\n/health + /metrics"),
        ("Operations", "Docker + K8s\nPrometheus + Grafana"),
    ]
    gap = 30
    card_width = (width - 2 * 35 - gap * (len(cards) - 1)) // len(cards)
    y0, y1 = 75, 335
    for index, (title, detail) in enumerate(cards):
        x0 = 35 + index * (card_width + gap)
        x1 = x0 + card_width
        draw.rounded_rectangle((x0, y0, x1, y1), radius=22, fill="#EAF1F7", outline="#2F6690", width=3)
        bbox = draw.textbbox((0, 0), title, font=font)
        draw.text((x0 + (card_width - (bbox[2] - bbox[0])) / 2, y0 + 32), title, fill="#17324D", font=font)
        for line_index, line in enumerate(detail.split("\n")):
            box = draw.textbbox((0, 0), line, font=small)
            draw.text((x0 + (card_width - (box[2] - box[0])) / 2, y0 + 110 + line_index * 42), line, fill="#53616F", font=small)
        if index < len(cards) - 1:
            arrow_x = x1 + 4
            center_y = (y0 + y1) // 2
            draw.line((arrow_x, center_y, arrow_x + gap - 8, center_y), fill="#2F6690", width=5)
            draw.polygon(
                [(arrow_x + gap - 8, center_y), (arrow_x + gap - 19, center_y - 8),
                 (arrow_x + gap - 19, center_y + 8)],
                fill="#2F6690",
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    style_run(run, size=8, color=GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build() -> Path:
    if not METRICS_PATH.exists():
        raise FileNotFoundError("Train the models first; artifacts/metrics.json is missing.")
    if not EDA_PATH.exists():
        raise FileNotFoundError("Run the EDA script first; artifacts/eda/eda_summary.json is missing.")
    metrics = json.loads(METRICS_PATH.read_text(encoding="utf-8"))
    eda = json.loads(EDA_PATH.read_text(encoding="utf-8"))
    verification = (
        json.loads(VERIFICATION_PATH.read_text(encoding="utf-8"))
        if VERIFICATION_PATH.exists()
        else {"ruff": "Not run", "pytest": "Not run", "training": "Not run"}
    )
    architecture_path = ROOT / "reports" / "figures" / "architecture.png"
    make_architecture(architecture_path)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = False

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("263238")
    normal.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.text = "MLOPS EXPERIMENTAL LEARNING  |  HEART DISEASE CLASSIFICATION"
    for run in header.runs:
        style_run(run, size=8, color=GRAY, bold=True)
    footer = section.footer.paragraphs[0]
    add_page_field(footer)

    # Page 1: cover, concise executive summary, and architecture.
    add_kicker(document, "Machine Learning Operations  |  Assignment 01  |  October 2026")
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.paragraph_format.space_before = Pt(45)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Heart Disease Classification")
    style_run(run, size=27, color="000000", bold=True)
    subtitle = add_paragraph(document, "An end-to-end MLOps workflow for the UCI Cleveland dataset", size=15, color=GRAY, after=22)
    subtitle.paragraph_format.keep_with_next = True
    add_table(
        document,
        ["Course", "Dataset", "Records", "Target"],
        [["MLOps AIMLCZG523", "UCI Heart Disease", "303", "Diagnosis > 0"]],
        widths=[1.65, 2.05, 1.0, 2.1],
        font_size=9,
    )
    add_heading(document, "Executive summary")
    add_paragraph(
        document,
        "This project builds a repeatable path from the supplied clinical data to a model-backed API. "
        "The source diagnosis is recoded as absence (0) versus presence (1). A single sklearn pipeline "
        "owns missing-value handling, feature scaling, categorical encoding, and classification so "
        "training and inference share the same transformations. Logistic Regression and Random Forest "
        "are tuned with stratified cross-validation; mean cross-validation ROC-AUC selects the model. "
        "The trained artifact is served by FastAPI, while MLflow records experiments and Prometheus "
        "metrics feed a provisioned Grafana dashboard.",
    )
    add_paragraph(
        document,
        "This report records measured model results and local runtime checks. Docker Compose served "
        "health and prediction requests, Prometheus reported the API target as UP, and Grafana received "
        "metrics. Docker Desktop Kind ran two API replicas, with health, prediction, and metrics "
        "verified through kubectl port-forward. No public cloud endpoint is used.",
        size=9.5,
        color=GRAY,
    )
    document.add_picture(str(architecture_path), width=Inches(6.85))
    add_paragraph(document, "Figure 1. Data, training, serving, and monitoring architecture.", size=8.5, color=GRAY, after=0)

    # Page 2: source and contract.
    add_page_break(document)
    add_page_title(document, "01  |  DATA ACQUISITION", "Dataset and data contract")
    add_paragraph(
        document,
        "The project uses the processed Cleveland subset in the UCI Heart Disease archive. The legacy "
        "file has 14 comma-separated fields without a header: 13 predictors and one diagnosis value. "
        "The supplied archive is retained at the workspace root; data/raw/processed.cleveland.data is "
        "the focused input used by the scripts. The download helper fetches the UCI ZIP and extracts "
        "only this named member, printing SHA-256 hashes for provenance.",
    )
    add_paragraph(
        document,
        "Acquisition was run from the UCI endpoint. The downloaded archive SHA-256 matched the supplied "
        "archive (b17cd273da9ce1caa4710fce80227ea454d4dbf9fcbc8e6a9121672751563adc); the extracted "
        "Cleveland file hash is a74b7efa387bc9d108d7d0115d831fe9b414b29ae7124f331b622b4efa0427c8.",
        size=9,
        color=GRAY,
    )
    add_table(
        document,
        ["Data check", "Observed", "Handling"],
        [
            ["Records", str(eda["source"]["raw_rows"]), "Keep usable diagnosis rows"],
            ["Predictors", "13 clinical fields", "Preserve source order and names"],
            ["Unknown values", str(eda["source"]["total_missing_values"]), "Keep missing until model fit"],
            ["Diagnosis", "Ordinal values 0-4", "Binary label: 0 vs 1-4"],
            ["Duplicates", str(eda["source"]["raw_duplicate_rows"]), "Remove exact duplicate rows"],
        ],
        widths=[1.45, 1.75, 3.6],
    )
    add_heading(document, "Target definition")
    add_paragraph(
        document,
        "The binary label follows the assignment framing and UCI's documented convention: diagnosis "
        "0 is absence, while diagnoses 1, 2, 3, or 4 represent presence. This transformation is explicit "
        "in src/data.py and tested. It avoids treating the original ordinal values as four separate "
        "classes when the project objective is binary classification.",
    )
    add_paragraph(
        document,
        "Predictors are age, sex, chest-pain type (cp), resting blood pressure, serum cholesterol, "
        "fasting blood sugar flag, resting ECG category, maximum heart rate, exercise angina, ST "
        "depression, slope, major vessels (ca), and thalassemia category (thal).",
    )
    add_heading(document, "Data quality and limits")
    missing = eda["source"]["missing_values_by_column"]
    missing_lines = ", ".join(f"{name}: {count}" for name, count in missing.items() if count)
    add_paragraph(
        document,
        f"Missingness is sparse in this subset ({missing_lines}). The analysis keeps these cells as nulls; "
        "it does not delete patient rows or fill them before data splitting. The preprocessing pipeline "
        "fits the imputers only on each training fold. This preserves the holdout and validation boundaries.",
    )
    add_paragraph(
        document,
        "The subset is small and historical. It is appropriate for an educational workflow demonstration, "
        "not for clinical use or for claims about present-day populations.",
        size=9.5,
        color=GRAY,
    )

    # Page 3: distribution and class balance.
    add_page_break(document)
    add_page_title(document, "02  |  EXPLORATORY ANALYSIS", "Feature distributions and class balance")
    add_paragraph(
        document,
        "The first EDA view summarizes all 13 predictors after the target mapping. Continuous features "
        "show broad patient variability, while coded fields such as chest-pain type and thalassemia type "
        "are discrete categories even though they are stored as numbers. That distinction informs the "
        "model pipeline: numerical measurements are scaled; encoded clinical codes are one-hot encoded.",
    )
    distribution_path = ROOT / "artifacts" / "eda" / "feature_distributions.png"
    if distribution_path.exists():
        distribution_paragraph = document.add_paragraph()
        distribution_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        distribution_paragraph.add_run().add_picture(str(distribution_path), width=Inches(5.45))
        add_paragraph(document, "Figure 2. Histograms for Cleveland clinical features.", size=8.5, color=GRAY, after=4)
    counts = eda["clean_data"]["target_distribution"]
    add_table(
        document,
        ["Target", "Meaning", "Patients", "Share"],
        [
            ["0", "No diagnosis", counts.get("0", 0), f"{counts.get('0', 0)/sum(counts.values()):.1%}"],
            ["1", "Diagnosis present", counts.get("1", 0), f"{counts.get('1', 0)/sum(counts.values()):.1%}"],
        ],
        widths=[0.8, 2.35, 1.2, 1.2],
    )
    add_paragraph(
        document,
        "The target is moderately imbalanced, so accuracy alone would be incomplete. Precision and recall "
        "are reported alongside ROC-AUC, and the candidate classifiers use class weighting to reduce the "
        "impact of unequal class frequencies.",
    )

    # Page 4: correlations and source missingness.
    add_page_break(document)
    add_page_title(document, "02  |  EXPLORATORY ANALYSIS", "Associations and missing source values")
    add_paragraph(
        document,
        "The correlation heatmap is descriptive: it shows linear associations among numeric encodings and "
        "the binary target. Ordinal category codes can create correlations that do not represent a true "
        "distance between categories, so the figure is used to orient feature review rather than as a "
        "feature-selection rule.",
    )
    heatmap_path = ROOT / "artifacts" / "eda" / "correlation_heatmap.png"
    if heatmap_path.exists():
        document.add_picture(str(heatmap_path), width=Inches(5.5))
        add_paragraph(document, "Figure 3. Pearson correlation matrix, including the binary target.", size=8.5, color=GRAY)
    balance_path = ROOT / "artifacts" / "eda" / "class_balance.png"
    missing_path = ROOT / "artifacts" / "eda" / "missingness.png"
    if balance_path.exists() and missing_path.exists():
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        for cell, image_path in zip(table.rows[0].cells, [balance_path, missing_path], strict=True):
            cell.width = Inches(3.2)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.paragraphs[0].add_run().add_picture(str(image_path), width=Inches(2.8))
    add_paragraph(document, "Figure 4. Class balance (left) and missingness before imputation (right).", size=8.5, color=GRAY)
    add_paragraph(
        document,
        "The raw file contains six missing feature values. Their source-level visibility is retained in "
        "the summary JSON and figure. All plots and counts are generated by scripts/eda.py rather than "
        "manually edited, so a new run recreates the report inputs.",
        size=9.5,
        after=0,
    )

    # Page 5: feature engineering and models.
    add_page_break(document)
    add_page_title(document, "03  |  FEATURE ENGINEERING", "Leakage-safe preprocessing and model selection")
    add_paragraph(
        document,
        "A ColumnTransformer routes five measurements (age, resting blood pressure, serum cholesterol, "
        "maximum heart rate, and ST depression) through median imputation and StandardScaler. Eight coded "
        "fields are treated as categorical and pass through most-frequent imputation followed by one-hot "
        "encoding. Unknown categories at inference are ignored safely. These transformations are placed "
        "inside the classifier pipeline, so each cross-validation fold learns its own imputation and "
        "scaling statistics.",
    )
    add_table(
        document,
        ["Candidate", "Role", "Search dimensions"],
        [
            ["Logistic Regression", "Interpretable linear baseline", "C; class_weight balanced"],
            ["Random Forest", "Nonlinear tree ensemble", "tree count, depth, leaf size; balanced subsampling"],
        ],
        widths=[1.6, 2.3, 3.0],
    )
    add_heading(document, "Evaluation protocol")
    add_paragraph(
        document,
        "A stratified 80/20 split is created with a fixed random seed. GridSearchCV tunes each candidate "
        "on the training portion using five stratified folds (three in the CI smoke run). The refit metric "
        "is ROC-AUC. Accuracy, precision, recall, and ROC-AUC means and standard deviations are retained "
        "for the best setting. The untouched holdout reports all four metrics once after tuning. The final "
        "serialized pipeline is then refit on all 303 rows for API use.",
    )
    add_paragraph(
        document,
        "The holdout score is a single estimate from a small dataset and should not be interpreted as a "
        "stable clinical performance guarantee. Cross-validation variability is therefore shown next to "
        "each mean.",
        size=9.5,
        color=GRAY,
    )

    # Page 6: actual results.
    add_page_break(document)
    add_page_title(document, "04  |  MODEL RESULTS", "Cross-validation and holdout performance")
    model_rows = []
    for item in metrics["models"]:
        cv = item["cv_metrics"]
        holdout = item["holdout_metrics"]
        model_rows.append([
            item["model_name"].replace("_", " ").title(),
            f"{cv['cv_accuracy_mean']:.3f} ± {cv['cv_accuracy_std']:.3f}",
            f"{cv['cv_precision_mean']:.3f} ± {cv['cv_precision_std']:.3f}",
            f"{cv['cv_recall_mean']:.3f} ± {cv['cv_recall_std']:.3f}",
            f"{cv['cv_roc_auc_mean']:.3f} ± {cv['cv_roc_auc_std']:.3f}",
        ])
    add_paragraph(
        document,
        "Table 1 reports training-set cross-validation scores for each selected parameter setting. "
        "The mean and standard deviation summarize fold-to-fold variation.",
    )
    add_table(
        document,
        ["Model", "Accuracy", "Precision", "Recall", "ROC-AUC"],
        model_rows,
        widths=[1.7, 1.2, 1.2, 1.0, 1.1],
        font_size=8.2,
    )
    holdout_rows = []
    for item in metrics["models"]:
        holdout = item["holdout_metrics"]
        holdout_rows.append([
            item["model_name"].replace("_", " ").title(),
            *(f"{holdout[key]:.3f}" for key in ["accuracy", "precision", "recall", "roc_auc"]),
        ])
    add_paragraph(document, "Table 2 reports the held-out 20% test results.")
    add_table(
        document,
        ["Model", "Accuracy", "Precision", "Recall", "ROC-AUC"],
        holdout_rows,
        widths=[1.7, 1.2, 1.2, 1.0, 1.1],
        font_size=8.2,
    )
    selected = metrics["selected_model"].replace("_", " ").title()
    selected_item = next(item for item in metrics["models"] if item["model_name"] == metrics["selected_model"])
    add_heading(document, "Selection")
    add_paragraph(
        document,
        f"{selected} is selected because it achieved the highest mean cross-validation ROC-AUC "
        f"({selected_item['cv_metrics']['cv_roc_auc_mean']:.3f}). Its margin over the other candidate "
        "is very small relative to fold-to-fold variation, so this run does not show a decisive model "
        "separation. The ranking rule is fixed before the holdout comparison; the holdout is not used "
        "to select hyperparameters. The selected parameters are recorded in metrics.json and the MLflow run.",
    )
    add_paragraph(
        document,
        "These values are experimental estimates on a small historical sample. They are not medical "
        "thresholds, and the API output must not guide patient care.",
        size=9.5,
        color=GRAY,
    )

    # Page 7: tracking and packaging.
    add_page_break(document)
    add_page_title(document, "05  |  EXPERIMENT TRACKING", "MLflow and reproducible artifacts")
    add_paragraph(
        document,
        "Each candidate run records the model name, chosen hyperparameters, cross-validation metrics, "
        "holdout metrics, fitted pipeline, and holdout ROC curve. The default tracking URI is a local "
        "file store under mlruns/. Run the UI from the repository root with mlflow ui to inspect the "
        "experiment. A separate metrics.json and model_comparison.csv make the result readable without "
        "opening the tracking interface. Model signatures use numeric input fields with a type that "
        "can represent missing values, matching the preprocessing contract.",
    )
    add_table(
        document,
        ["Artifact", "Contents", "Use"],
        [
            ["heart_disease_model.joblib", "Preprocessor plus selected estimator", "API inference"],
            ["model_metadata.json", "Model family, parameters, features, time", "Model identity"],
            ["metrics.json", "Data snapshot, selection, CV and holdout results", "Reporting and audit"],
            ["model_comparison.csv", "One row per candidate", "Quick comparison"],
            ["*_roc.png", "Holdout ROC curve per model", "Visual assessment"],
            ["MLflow runs", "Parameters, metrics, model and plots", "Experiment history"],
        ],
        widths=[2.0, 2.75, 1.65],
    )
    add_heading(document, "Reproduction controls")
    add_bullet(document, "The raw source file is headerless and parsed through a shared schema.")
    add_bullet(document, "A fixed random seed controls data splitting, folds, and estimators.")
    add_bullet(document, "Preprocessing is fitted within cross-validation and serialized with the model.")
    add_bullet(document, "requirements.txt and the CI workflow define the execution environment.")
    add_paragraph(
        document,
        "MLflow organizes runs, records parameters and metrics, and stores output artifacts as documented "
        "in the official tracking API reference. The current setup is local; a shared server and object "
        "store would be needed for team-wide or cloud tracking.",
    )

    # Page 8: API contract and example.
    add_page_break(document)
    add_page_title(document, "06  |  MODEL SERVING", "FastAPI request and response contract")
    add_paragraph(
        document,
        "The API loads the packaged sklearn pipeline from MODEL_PATH and accepts one JSON object containing "
        "the 13 named predictors. Pydantic validates required fields and plausible input ranges before "
        "inference. Optional ca and thal values may be null; the trained pipeline handles them. The response "
        "returns a 0/1 prediction, a readable class label, and the model's estimated positive-class "
        "probability as confidence. This probability has not been calibrated or validated for clinical use.",
    )
    add_heading(document, "Example request")
    code = document.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(8)
    sample = (
        '{\n  "age": 63, "sex": 1, "cp": 1, "trestbps": 145,\n'
        '  "chol": 233, "fbs": 1, "restecg": 2, "thalach": 150,\n'
        '  "exang": 0, "oldpeak": 2.3, "slope": 3, "ca": 0, "thal": 6\n}'
    )
    run = code.add_run(sample)
    style_run(run, size=9, color="263238", name="Consolas")
    add_paragraph(document, "POST /predict   |   GET /health   |   GET /metrics", size=10, color=BLUE, bold=True)
    add_table(
        document,
        ["Endpoint", "Purpose", "Expected behavior"],
        [
            ["/predict", "One prediction", "200 with prediction, label, confidence, version"],
            ["/health", "Readiness", "Model load state; 503 when artifact unavailable"],
            ["/metrics", "Monitoring", "Prometheus request, latency, and class counters"],
        ],
        widths=[1.3, 1.55, 3.7],
    )
    add_heading(document, "Privacy and logging")
    add_paragraph(
        document,
        "Request logs capture the method, route, HTTP status, and elapsed time. They omit the request body "
        "so patient-feature values are not written to application logs. This reduces one operational "
        "exposure but does not replace access control, transport security, retention rules, or a privacy "
        "review for real health data.",
    )
    add_paragraph(
        document,
        "The local Python process and Docker Compose smoke checks returned /health status ok and valid "
        "/predict responses for samples/predict.json. The Docker container also returned HTTP 200 from "
        "/metrics. Kubernetes health, prediction, and metrics endpoints were verified separately through "
        "kubectl port-forward.",
        size=9.5,
        color=GRAY,
    )

    # Page 9: test and CI.
    add_page_break(document)
    add_page_title(document, "07  |  CI AND QUALITY", "Automated checks and local verification")
    add_paragraph(
        document,
        "The GitHub Actions workflow installs the requirements file, runs Ruff linting, executes pytest, "
        "trains the models in a smaller CI search mode, and uploads the model, metrics, plots, and local "
        "MLflow store as a workflow artifact. Each command is a required step, so a lint, test, or training "
        "failure marks the run as failed.",
    )
    add_table(
        document,
        ["Check", "Coverage", "Local status"],
        [
            ["Ruff", "Imports, style, common correctness errors", verification.get("ruff", "Not run")],
            ["Pytest", "Data contract, recoding, pipeline, API schema", verification.get("pytest", "Not run")],
            ["Training smoke run", "Search, metrics, plots, serialized model", verification.get("training", "Not run")],
            ["Workflow artifact", "Retain training outputs per CI run", verification.get("github_ci", "Not run")],
        ],
        widths=[1.5, 3.1, 1.95],
    )
    add_heading(document, "Reproduction commands")
    for command in [
        "ruff check src app scripts tests",
        "pytest",
        "python scripts/train.py --fast",
    ]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.15)
        paragraph.paragraph_format.space_after = Pt(3)
        style_run(paragraph.add_run(command), size=9, color="263238", name="Consolas")
    add_paragraph(
        document,
        f"Current GitHub Actions status: {verification.get('github_ci', 'Not run')}. "
        "Capture the completed first run page as "
        "docs/screenshots/github-actions-ci-passed.png before submission.",
    )

    # Page 10: deployment, monitoring, scope, and references.
    add_page_break(document)
    add_page_title(document, "08  |  DEPLOYMENT AND OPERATIONS", "Container, Kubernetes, and monitoring")
    add_paragraph(
        document,
        "The Docker image packages the API and serialized model. Docker Compose starts the API with "
        "Prometheus and Grafana; the API exports request totals, latency histograms, and prediction counts. "
        "A provisioned dashboard shows request rate, p95 latency, class counts, and 5xx ratio. The "
        "Kubernetes Deployment specifies two replicas, readiness and liveness probes, resource bounds, "
        "and scrape annotations; a LoadBalancer Service exposes port 80.",
    )
    add_table(
        document,
        ["Evidence item", "What is prepared", "What remains"],
        [
            ["API container", "Dockerfile; health, prediction, and metrics requests", "Passed locally"],
            ["Observability", "Prometheus target UP; Grafana dashboard", "Passed locally"],
            ["Kubernetes", "Two ready replicas; endpoints via port-forward", "Passed locally"],
            ["Public deployment", "No cloud-specific resource is assumed", "Not performed; no public URL"],
            ["Video and screenshots", "Three runtime screenshots in docs/screenshots", "Short demo video remains"],
        ],
        widths=[1.35, 2.7, 2.5],
        font_size=8.1,
    )
    add_heading(document, "Conclusion and limitations")
    add_paragraph(
        document,
        "Docker Compose and the local Docker Desktop Kind cluster have both been exercised. The Kubernetes "
        "Deployment reached two ready replicas, and /health, /predict, and /metrics were verified through "
        "port-forwarding. Prometheus, Grafana, and Kubernetes screenshots are included in the repository. "
        f"The current GitHub Actions result is {verification.get('github_ci', 'not yet run')}. No public "
        "cloud endpoint is claimed; a short demo video and the course submission remain for the student. "
        "The dataset has only 303 historical records, so model metrics should not be used for diagnosis "
        "or health decisions.",
        size=9.7,
    )
    add_heading(document, "References", level=2)
    add_paragraph(
        document,
        "UCI Machine Learning Repository. Heart Disease dataset, DOI 10.24432/C52P4X. "
        "https://archive.ics.uci.edu/dataset/45/heart+disease",
        size=8.2,
        after=3,
    )
    add_paragraph(
        document,
        "MLflow. Tracking API documentation. https://mlflow.org/docs/latest/tracking/tracking-api/",
        size=8.2,
        after=3,
    )
    add_paragraph(
        document,
        "scikit-learn. ColumnTransformer and model evaluation documentation. "
        "https://scikit-learn.org/stable/modules/generated/sklearn.compose.ColumnTransformer.html",
        size=8.2,
        after=0,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "Heart Disease Classification MLOps Assignment Report"
    document.core_properties.subject = "End-to-end MLOps assignment"
    document.core_properties.author = "MLOps Assignment"
    document.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())

